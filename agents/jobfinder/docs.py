"""docx/pdf helpers."""
from __future__ import annotations

import copy
import io
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"


# extraction
def _iter_paragraphs(doc):
    """Body paragraphs, then table cells (lots of resume templates are tables)."""
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:  # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                yield from cell.paragraphs


def docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in _iter_paragraphs(doc) if p.text.strip())


def pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def any_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        return docx_text(data)
    if name.endswith(".pdf"):
        return pdf_text(data)
    return data.decode("utf-8", errors="ignore")


# tailoring
def numbered_paragraphs(data: bytes) -> list[dict]:
    """[{idx, text}] for non-empty paragraphs. The tailor edits by idx."""
    doc = Document(io.BytesIO(data))
    return [{"idx": i, "text": p.text} for i, p in enumerate(_iter_paragraphs(doc)) if p.text.strip()]


def _set_text(paragraph, new_text: str) -> None:
    """Replace the text but keep formatting. Keeps a leading label run like a bold "Skills:"."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return
    label = runs[0].text
    if len(runs) > 1 and label.strip() and new_text.startswith(label):
        runs[1].text = new_text[len(label):]
        for r in runs[2:]:
            r.text = ""
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def apply_edits(data: bytes, edits: list[dict], inserts: list[dict]) -> bytes:
    """Apply edits [{idx, new_text}] and inserts [{after_idx, text}] (inserts copy the paragraph's style)."""
    doc = Document(io.BytesIO(data))
    paragraphs = list(_iter_paragraphs(doc))
    for e in edits:
        i = int(e.get("idx", -1))
        if 0 <= i < len(paragraphs) and e.get("new_text", "").strip():
            _set_text(paragraphs[i], e["new_text"].strip())
    # bottom-up so indices stay valid
    for ins in sorted(inserts, key=lambda x: int(x.get("after_idx", -1)), reverse=True):
        i = int(ins.get("after_idx", -1))
        if 0 <= i < len(paragraphs) and ins.get("text", "").strip():
            src = paragraphs[i]
            new_p = copy.deepcopy(src._p)
            src._p.addnext(new_p)
            _set_text(Paragraph(new_p, src._parent), ins["text"].strip())
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# cover letter
def cover_letter_docx(header: dict, letter: dict) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(64)
        section.top_margin = section.bottom_margin = Pt(56)

    name = doc.add_paragraph()
    run = name.add_run(header.get("name") or "")
    run.bold = True
    run.font.size = Pt(16)
    contact = " | ".join(x for x in [header.get("email"), header.get("phone"), header.get("location"), header.get("link")] if x and x != "NA")
    if contact:
        doc.add_paragraph(contact).runs[0].font.size = Pt(9.5)
    doc.add_paragraph(letter.get("date", ""))
    recipient = "\n".join(x for x in [letter.get("recipient"), letter.get("company"), letter.get("company_location")] if x and x != "NA")
    if recipient:
        doc.add_paragraph(recipient)
    if letter.get("subject"):
        doc.add_paragraph().add_run(letter["subject"]).bold = True
    doc.add_paragraph(letter.get("greeting", "Dear Hiring Manager,"))
    for para in letter.get("paragraphs", []):
        p = doc.add_paragraph(para)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
    doc.add_paragraph(letter.get("closing", "Sincerely,"))
    doc.add_paragraph(header.get("name") or "")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# pdf
def _soffice() -> str | None:
    for cand in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        path = shutil.which(cand) or (cand if Path(cand).exists() else None)
        if path:
            return path
    return None


def docx_to_pdf(data: bytes, stem: str) -> bytes | None:
    """LibreOffice conversion, None if soffice isn't installed."""
    exe = _soffice()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / f"{stem}.docx"
        src.write_bytes(data)
        subprocess.run(
            [exe, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
            check=True, capture_output=True, timeout=180,
        )
        pdf = Path(tmp) / f"{stem}.pdf"
        return pdf.read_bytes() if pdf.exists() else None
